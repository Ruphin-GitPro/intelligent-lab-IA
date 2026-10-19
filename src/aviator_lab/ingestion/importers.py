"""Import de résultats CSV et Excel, avec rapport explicite par ligne."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from aviator_lab.storage import DuplicateBetError, DuplicateRoundError, RoundRepository


class ImportFileError(ValueError):
    """Fichier illisible ou format non pris en charge."""


@dataclass(frozen=True)
class ImportIssue:
    """Erreur associée à une ligne du fichier d'entrée."""

    row_number: int
    message: str


@dataclass(frozen=True)
class ImportReport:
    """Résumé sans perte silencieuse des lignes rejetées."""

    imported: int
    rejected: tuple[ImportIssue, ...]
    skipped_empty: int


def _read_rows(path: Path) -> tuple[list[str], Iterable[tuple[int, dict[str, Any]]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            stream = path.open("r", encoding="utf-8-sig", newline="")
            reader = csv.reader(stream)
            raw_headers = next(reader, None)
            if not raw_headers:
                stream.close()
                raise ImportFileError("Le CSV est vide ou ne contient pas d'en-tête.")
            headers = [str(name).strip().lower() for name in raw_headers]

            def csv_rows():
                try:
                    for values_row in reader:
                        row = dict(zip(headers, values_row))
                        if len(values_row) > len(headers):
                            row["__extra_fields__"] = values_row[len(headers) :]
                        yield reader.line_num, row
                finally:
                    stream.close()

            return headers, csv_rows()
        except OSError as exc:
            raise ImportFileError(f"Impossible de lire le fichier : {exc}") from exc

    if suffix in {".xlsx", ".xlsm"}:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            raw_headers = next(values, None)
            if not raw_headers:
                workbook.close()
                raise ImportFileError("La feuille Excel est vide ou ne contient pas d'en-tête.")
            headers = [str(value).strip().lower() if value is not None else "" for value in raw_headers]

            def excel_rows():
                try:
                    for row_number, values_row in enumerate(values, start=2):
                        yield row_number, dict(zip(headers, values_row))
                finally:
                    workbook.close()

            return headers, excel_rows()
        except OSError as exc:
            raise ImportFileError(f"Impossible de lire le fichier Excel : {exc}") from exc

    raise ImportFileError("Format non pris en charge. Utilisez un fichier .csv ou .xlsx.")


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_datetime(value: Any) -> datetime | None:
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError("date invalide (format ISO 8601 attendu)") from exc
    raise ValueError("date invalide (format ISO 8601 attendu)")


def import_rounds(
    file_path: str | Path,
    repository: RoundRepository,
    *,
    default_origin: str = "real",
) -> ImportReport:
    """Importe les lignes valides et retourne les problèmes avec leur numéro de ligne.

    En-têtes requis : ``multiplier``. En-têtes facultatifs : ``external_id``,
    ``played_at`` (ISO 8601) et ``origin``. Chaque insertion est indépendante;
    une ligne invalide n'efface pas les autres et est comptabilisée dans le rapport.
    """
    path = Path(file_path).expanduser()
    headers, rows = _read_rows(path)
    if "multiplier" not in headers:
        raise ImportFileError("Colonne obligatoire absente : multiplier.")

    imported = 0
    skipped_empty = 0
    rejected: list[ImportIssue] = []
    source = f"import:{path.name}"
    for row_number, row in rows:
        if all(_blank(value) for value in row.values()):
            skipped_empty += 1
            continue
        try:
            if "__extra_fields__" in row:
                raise ValueError("la ligne contient plus de champs que l'en-tête")
            played_at = _parse_datetime(row.get("played_at"))
            origin_value = row.get("origin")
            origin = default_origin if _blank(origin_value) else str(origin_value).strip().lower()
            external_id_value = row.get("external_id")
            external_id = None if _blank(external_id_value) else str(external_id_value).strip()
            repository.add_round(
                row.get("multiplier"),
                source=source,
                origin=origin,
                external_id=external_id,
                played_at=played_at,
            )
            imported += 1
        except (ValueError, TypeError, DuplicateRoundError) as exc:
            rejected.append(ImportIssue(row_number, str(exc)))

    return ImportReport(imported, tuple(rejected), skipped_empty)


def import_bets(
    file_path: str | Path,
    repository: RoundRepository,
    *,
    default_origin: str = "real",
) -> ImportReport:
    """Importe des paris personnels sans supposer un retour depuis le cashout.

    En-têtes possibles : ``external_id``, ``round_id``, ``bet_at``,
    ``stake_amount``, ``cashout_multiplier``, ``returned_amount``, ``net_amount``
    ``currency`` et ``origin``. Les champs financiers sont conservés comme manquants s'ils ne
    sont pas fournis. Le net n'est calculé que si mise et retour sont tous deux
    renseignés, par la règle arithmétique ``retour - mise``.
    """
    path = Path(file_path).expanduser()
    headers, rows = _read_rows(path)
    supported = {
        "external_id",
        "round_id",
        "bet_at",
        "stake_amount",
        "cashout_multiplier",
        "returned_amount",
        "net_amount",
        "currency",
        "origin",
    }
    unknown = set(headers) - supported
    if unknown:
        raise ImportFileError(f"Colonnes de paris inconnues : {', '.join(sorted(unknown))}.")

    imported = 0
    skipped_empty = 0
    rejected: list[ImportIssue] = []
    source = f"import:{path.name}"
    for row_number, row in rows:
        if all(_blank(value) for value in row.values()):
            skipped_empty += 1
            continue
        try:
            if "__extra_fields__" in row:
                raise ValueError("la ligne contient plus de champs que l'en-tête")
            round_value = row.get("round_id")
            round_id = None
            if not _blank(round_value):
                try:
                    round_decimal = Decimal(str(round_value).strip())
                except (InvalidOperation, ValueError) as exc:
                    raise ValueError("identifiant de manche invalide") from exc
                if not round_decimal.is_finite() or round_decimal != round_decimal.to_integral_value():
                    raise ValueError("identifiant de manche invalide")
                round_id = int(round_decimal)
            origin_value = row.get("origin")
            origin = default_origin if _blank(origin_value) else str(origin_value).strip().lower()
            external_id_value = row.get("external_id")
            external_id = None if _blank(external_id_value) else str(external_id_value).strip()
            repository.add_bet(
                source=source,
                origin=origin,
                external_id=external_id,
                round_id=round_id,
                bet_at=_parse_datetime(row.get("bet_at")),
                stake_amount=row.get("stake_amount"),
                cashout_multiplier=row.get("cashout_multiplier"),
                returned_amount=row.get("returned_amount"),
                net_amount=row.get("net_amount"),
                currency=None if _blank(row.get("currency")) else str(row["currency"]),
            )
            imported += 1
        except (ValueError, TypeError, DuplicateBetError) as exc:
            rejected.append(ImportIssue(row_number, str(exc)))
    return ImportReport(imported, tuple(rejected), skipped_empty)
